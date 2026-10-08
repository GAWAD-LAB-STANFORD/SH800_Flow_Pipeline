#!/usr/bin/env python3
"""
Shared lab-document builder.

Renders a JSON "doc spec" into a uniformly formatted .docx so that every lab
skill (results report, protocol, brainstorm, schedule) produces documents that
look identical for the PI to read.

Usage:
    python build_docx.py spec.json output.docx

The spec schema is documented in house_style.md (same folder). Keep this file
BYTE-IDENTICAL across all lab skills — that shared rendering is what makes the
outputs uniform. If you change styling, change it here and copy to every skill.
"""

import sys
import json
import os

# --- self-bootstrap python-docx so this runs in a fresh Cowork session ---
try:
    import docx  # noqa
except ImportError:
    import subprocess
    subprocess.run(
        [sys.executable, "-m", "pip", "install", "--break-system-packages", "-q", "python-docx"],
        check=True,
    )

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- House style constants (the single source of truth for look & feel) ----
ACCENT = "8C1515"          # Stanford cardinal — heading + rule color
INK = "1A1A1A"             # body text
MUTED = "5A5A5A"           # captions, footer
LIGHT_FILL = "F3EEEE"      # metadata / table header shading
BODY_FONT = "Calibri"
BODY_SIZE = 10.5
H1_SIZE = 13
TITLE_SIZE = 18


def _set_font(run, name=BODY_FONT, size=BODY_SIZE, bold=False, color=INK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


def _thin_borders(table, color=ACCENT):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    for edge in ("left", "right", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


def _bottom_rule(paragraph, color=ACCENT, size=12):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def add_title(doc, spec):
    # small doc-type kicker
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    r = kicker.add_run((spec.get("doc_type", "Lab Document")).upper())
    _set_font(r, size=9, bold=True, color=ACCENT)
    r.font.name = BODY_FONT
    # spaced-out letters for the kicker
    rPr = r._element.get_or_add_rPr()
    spc = OxmlElement("w:spacing")
    spc.set(qn("w:val"), "40")
    rPr.append(spc)

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(6)
    r = t.add_run(spec.get("title", "Untitled"))
    _set_font(r, size=TITLE_SIZE, bold=True, color=INK)
    _bottom_rule(t)


def add_meta_block(doc, meta):
    if not meta:
        return
    items = list(meta.items())
    tbl = doc.add_table(rows=len(items), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _no_borders(tbl)
    _cell_margins(tbl)
    tbl.columns[0].width = Inches(1.6)
    tbl.columns[1].width = Inches(5.2)
    for i, (k, v) in enumerate(items):
        c0, c1 = tbl.rows[i].cells
        c0.width = Inches(1.6)
        c1.width = Inches(5.2)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(str(k))
        _set_font(r0, size=9.5, bold=True, color=MUTED)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(str(v))
        _set_font(r1, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _set_font(r, size=H1_SIZE, bold=True, color=ACCENT)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    _set_font(r)


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(str(it))
        _set_font(r)


def add_keyvalue(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=2)
    _thin_borders(tbl)
    _cell_margins(tbl)
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(4.6)
    for i, (k, v) in enumerate(rows):
        c0, c1 = tbl.rows[i].cells
        c0.width = Inches(2.2)
        c1.width = Inches(4.6)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(str(k))
        _set_font(r0, size=10, bold=True, color=INK)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(str(v))
        _set_font(r1, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, columns, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(columns))
    _thin_borders(tbl)
    _cell_margins(tbl)
    hdr = tbl.rows[0].cells
    for j, col in enumerate(columns):
        _shade(hdr[j], LIGHT_FILL)
        p = hdr[j].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(col))
        _set_font(r, size=9.5, bold=True, color=ACCENT)
    for i, row in enumerate(rows):
        cells = tbl.rows[i + 1].cells
        for j, val in enumerate(row):
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run("" if val is None else str(val))
            _set_font(r, size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc, path, caption=None, width_in=5.6):
    if path and os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        try:
            p.add_run().add_picture(path, width=Inches(width_in))
        except Exception as e:
            r = p.add_run(f"[Could not embed image {path}: {e}]")
            _set_font(r, size=9, italic=True, color=MUTED)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"[Figure placeholder — attach: {path or 'image'}]")
        _set_font(r, size=9, italic=True, color=MUTED)
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(8)
        r = cp.add_run(caption)
        _set_font(r, size=9, italic=True, color=MUTED)


def add_callout(doc, text, label="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    _cell_margins(tbl, top=100, bottom=100, left=140, right=140)
    cell = tbl.rows[0].cells[0]
    _shade(cell, LIGHT_FILL)
    # left accent bar
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), ACCENT)
    borders.append(left)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}  ")
    _set_font(r, size=9, bold=True, color=ACCENT)
    r2 = p.add_run(text)
    _set_font(r2, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_blocks(doc, blocks):
    for b in blocks:
        bt = b.get("type")
        if bt == "paragraph":
            add_paragraph(doc, b.get("text", ""))
        elif bt == "bullets":
            add_bullets(doc, b.get("items", []))
        elif bt == "numbered":
            add_bullets(doc, b.get("items", []), numbered=True)
        elif bt == "keyvalue":
            add_keyvalue(doc, b.get("rows", []))
        elif bt == "table":
            add_table(doc, b.get("columns", []), b.get("rows", []))
        elif bt == "figure":
            add_figure(doc, b.get("path"), b.get("caption"))
        elif bt == "callout":
            add_callout(doc, b.get("text", ""), b.get("label", "NOTE"))
        else:
            add_paragraph(doc, b.get("text", ""))


def set_base_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(INK)


def add_footer(doc, spec):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bits = [spec.get("footer", "")]
    text = "   ·   ".join([b for b in bits if b])
    if not text:
        text = f"{spec.get('doc_type','Lab Document')}"
    r = p.add_run(text)
    _set_font(r, size=8, color=MUTED)


def build(spec_path, out_path):
    with open(spec_path) as f:
        spec = json.load(f)

    doc = Document()
    set_base_style(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    add_title(doc, spec)
    add_meta_block(doc, spec.get("meta", {}))

    for sec in spec.get("sections", []):
        if sec.get("heading"):
            add_heading(doc, sec["heading"])
        render_blocks(doc, sec.get("blocks", []))

    add_footer(doc, spec)
    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_docx.py spec.json output.docx")
        sys.exit(1)
    build(sys.argv[1], sys.argv[2])
